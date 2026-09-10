from pathlib import Path

from openpyxl import Workbook

from scripts.dataverse_role_import import read_workbook_roles


def test_read_workbook_roles_keeps_highest_duplicate_privilege_depth(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    roles = workbook.active
    roles.title = "Roles"
    roles.append(
        [
            "Role Index",
            "Role Name",
            "Role Name (Exact XML)",
            "Role ID",
        ]
    )
    roles.append([1, "Display Name", "Exact Role", "source-role-id"])

    privileges = workbook.create_sheet("Role_Privileges")
    privileges.append(
        [
            "Role Index",
            "Role Name",
            "Role ID",
            "Source File",
            "Privilege Name",
            "Action",
            "Target/Table",
            "Access Level",
            "Access Scope Meaning",
        ]
    )
    privileges.append(
        [
            1,
            "Exact Role",
            "source-role-id",
            "a.xml",
            "prvReadAccount",
            "Read",
            "account",
            "Basic",
            "User",
        ]
    )
    privileges.append(
        [
            1,
            "Exact Role",
            "source-role-id",
            "a.xml",
            "prvReadAccount",
            "Read",
            "account",
            "Global",
            "Organization",
        ]
    )
    privileges.append(
        [
            1,
            "Exact Role",
            "source-role-id",
            "a.xml",
            "prvWriteAccount",
            "Write",
            "account",
            "Local",
            "Business Unit",
        ]
    )

    path = tmp_path / "roles.xlsx"
    workbook.save(path)

    result = read_workbook_roles(path)

    assert result == [
        {
            "source_role_index": 1,
            "source_role_indices": [1],
            "source_role_id": "source-role-id",
            "source_role_ids": ["source-role-id"],
            "name": "Exact Role",
            "privileges": {
                "prvReadAccount": "Global",
                "prvWriteAccount": "Local",
            },
        }
    ]


def test_read_workbook_roles_skips_unknown_depth(tmp_path: Path) -> None:
    workbook = Workbook()
    roles = workbook.active
    roles.title = "Roles"
    roles.append(["Role Index", "Role Name", "Role Name (Exact XML)", "Role ID"])
    roles.append([1, "Role", "Role", "role-id"])

    privileges = workbook.create_sheet("Role_Privileges")
    privileges.append(
        [
            "Role Index",
            "Role Name",
            "Role ID",
            "Source File",
            "Privilege Name",
            "Action",
            "Target/Table",
            "Access Level",
            "Access Scope Meaning",
        ]
    )
    privileges.append(
        [
            1,
            "Role",
            "role-id",
            "a.xml",
            "prvReadAccount",
            "Read",
            "account",
            "RecordFilter",
            "Custom",
        ]
    )

    path = tmp_path / "roles.xlsx"
    workbook.save(path)

    result = read_workbook_roles(path)

    assert result[0]["privileges"] == {}


def test_read_workbook_roles_merges_duplicate_exact_names(tmp_path: Path) -> None:
    workbook = Workbook()
    roles = workbook.active
    roles.title = "Roles"
    roles.append(["Role Index", "Role Name", "Role Name (Exact XML)", "Role ID"])
    roles.append([1, "Role A", "Shared Role", "role-a"])
    roles.append([2, "Role B", "Shared Role", "role-b"])

    privileges = workbook.create_sheet("Role_Privileges")
    privileges.append(
        [
            "Role Index",
            "Role Name",
            "Role ID",
            "Source File",
            "Privilege Name",
            "Action",
            "Target/Table",
            "Access Level",
            "Access Scope Meaning",
        ]
    )
    privileges.append(
        [
            1,
            "Role A",
            "role-a",
            "a.xml",
            "prvReadAccount",
            "Read",
            "account",
            "Basic",
            "User",
        ]
    )
    privileges.append(
        [
            2,
            "Role B",
            "role-b",
            "b.xml",
            "prvReadAccount",
            "Read",
            "account",
            "Deep",
            "Parent Child",
        ]
    )
    privileges.append(
        [
            2,
            "Role B",
            "role-b",
            "b.xml",
            "prvReadContact",
            "Read",
            "contact",
            "Local",
            "Business Unit",
        ]
    )

    path = tmp_path / "roles.xlsx"
    workbook.save(path)

    result = read_workbook_roles(path)

    assert len(result) == 1
    assert result[0]["source_role_indices"] == [1, 2]
    assert result[0]["source_role_ids"] == ["role-a", "role-b"]
    assert result[0]["privileges"] == {
        "prvReadAccount": "Deep",
        "prvReadContact": "Local",
    }
