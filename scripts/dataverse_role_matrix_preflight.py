"""Validate a Dataverse security-role workbook without contacting services."""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_COLUMNS = {
    "Roles": {"Role Name (Exact XML)", "Role ID"},
    "Role_Privileges": {
        "Role Name",
        "Privilege Name",
        "Action",
        "Target/Table",
        "Access Level",
    },
}
ENTITLEMENT_DIMENSIONS = {
    "business_units": {"business_units", "businessunits", "business units"},
    "users": {"users", "systemusers", "system users"},
    "teams": {"teams", "team_memberships", "team memberships"},
    "field_security_profiles": {
        "field_security_profiles",
        "field security profiles",
    },
    "poa": {"poa", "principalobjectaccess"},
    "poaa": {"poaa", "principalobjectattributeaccess"},
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create a read-only readiness report for a Dataverse role matrix."
    )
    parser.add_argument(
        "--workbook", default="Dataverse_Security_Roles_Consolidated.xlsx"
    )
    parser.add_argument(
        "--output", default="reports/dataverse-role-matrix-preflight.json"
    )
    parser.add_argument("--expected-role-count", type=int)
    return parser.parse_args()


def rows_as_dicts(sheet) -> list[dict[str, Any]]:
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    return [
        dict(zip(headers, values))
        for values in iterator
        if any(value is not None for value in values)
    ]


def validate_columns(workbook) -> None:
    for sheet_name, required in REQUIRED_COLUMNS.items():
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Workbook is missing required sheet '{sheet_name}'.")
        headers = {
            cell.value
            for cell in next(workbook[sheet_name].iter_rows(min_row=1, max_row=1))
        }
        missing = required.difference(headers)
        if missing:
            raise ValueError(
                f"Sheet '{sheet_name}' is missing columns: {sorted(missing)}"
            )


def build_report(path: Path, expected_role_count: int | None) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    validate_columns(workbook)
    roles = rows_as_dicts(workbook["Roles"])
    privileges = rows_as_dicts(workbook["Role_Privileges"])

    role_names = [row.get("Role Name (Exact XML)") for row in roles]
    role_ids = [row.get("Role ID") for row in roles if row.get("Role ID")]
    names_by_role_id: dict[str, set[str]] = defaultdict(set)
    for row in roles:
        if row.get("Role ID") and row.get("Role Name (Exact XML)"):
            names_by_role_id[str(row["Role ID"])].add(str(row["Role Name (Exact XML)"]))

    action_counts = Counter(
        str(row.get("Action") or "Unknown").strip() for row in privileges
    )
    read_rows = [
        row
        for row in privileges
        if str(row.get("Action") or "").strip().casefold() == "read"
    ]
    depth_counts = Counter(
        str(row.get("Access Level") or "Other/Blank").strip() or "Other/Blank"
        for row in read_rows
    )
    sheet_names = {name.casefold() for name in workbook.sheetnames}
    dimensions_present = {
        dimension: bool(sheet_names.intersection(aliases))
        for dimension, aliases in ENTITLEMENT_DIMENSIONS.items()
    }
    role_count = len(roles)
    expected_count_matches = (
        None if expected_role_count is None else role_count == expected_role_count
    )

    findings = []
    if expected_role_count is not None and not expected_count_matches:
        findings.append(
            f"Workbook contains {role_count:,} role rows; expected "
            f"{expected_role_count:,}."
        )
    missing_dimensions = [
        name for name, present in dimensions_present.items() if not present
    ]
    if missing_dimensions:
        findings.append(
            "Workbook cannot prove effective entitlements because it lacks: "
            + ", ".join(missing_dimensions)
            + "."
        )

    return {
        "workbook": str(path),
        "status": "incomplete" if findings else "structurally-valid",
        "fabric_sync_ready": not findings,
        "expected_role_count": expected_role_count,
        "expected_role_count_matches": expected_count_matches,
        "sheets": workbook.sheetnames,
        "roles": {
            "rows": role_count,
            "unique_names": len(set(role_names)),
            "unique_ids": len(set(role_ids)),
            "duplicate_name_rows": sum(
                count - 1 for count in Counter(role_names).values() if count > 1
            ),
            "ids_used_by_multiple_names": sum(
                len(names) > 1 for names in names_by_role_id.values()
            ),
        },
        "privileges": {
            "rows": len(privileges),
            "actions": dict(sorted(action_counts.items())),
            "read_rows": len(read_rows),
            "read_depths": dict(sorted(depth_counts.items())),
            "read_targets": len(
                {
                    row.get("Target/Table")
                    for row in read_rows
                    if row.get("Target/Table")
                }
            ),
        },
        "entitlement_dimensions_present": dimensions_present,
        "findings": findings,
    }


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.workbook)
    report = build_report(workbook_path, args.expected_role_count)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    print(json.dumps(report, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
