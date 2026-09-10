"""Create an aggregate Dataverse environment comparison workbook."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
GOOD_FILL = PatternFill("solid", fgColor="E2F0D9")
HEADER_FONT = Font(name="Arial", color="FFFFFF", bold=True)
TITLE_FONT = Font(name="Arial", size=16, bold=True, color="1F1F1F")
BODY_FONT = Font(name="Arial", size=10, color="000000")
INPUT_FONT = Font(name="Arial", size=10, color="0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="000000")


def _flatten(value: dict[str, Any], prefix: str = "") -> list[tuple[str, Any]]:
    rows: list[tuple[str, Any]] = []
    for key, child in value.items():
        path = f"{prefix}.{key}" if prefix else key
        if isinstance(child, dict):
            rows.extend(_flatten(child, path))
        else:
            rows.append((path, child))
    return rows


def _client_workbook_summary(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    role_rows = list(workbook["Roles"].iter_rows(min_row=2, values_only=True))
    grant_rows = list(
        workbook["Role_Privileges"].iter_rows(min_row=2, values_only=True)
    )
    read_rows = [
        row for row in grant_rows if str(row[5] or "").strip().lower() == "read"
    ]
    all_depths = Counter(str(row[7] or "Blank") for row in grant_rows)
    read_depths = Counter(str(row[7] or "Blank") for row in read_rows)

    return {
        "exported_roles": len(role_rows),
        "distinct_role_ids": len({row[3] for row in role_rows if row[3]}),
        "unique_role_names": len({row[2] for row in role_rows if row[2]}),
        "deprecated_roles": sum(
            str(row[13]).strip().lower() in {"true", "yes", "1"} for row in role_rows
        ),
        "total_privilege_grants": len(grant_rows),
        "unique_privilege_names": len({row[4] for row in grant_rows if row[4]}),
        "unique_targets_all_actions": len({row[6] for row in grant_rows if row[6]}),
        "read_grants": len(read_rows),
        "read_unique_privileges": len({row[4] for row in read_rows if row[4]}),
        "read_unique_targets": len({row[6] for row in read_rows if row[6]}),
        "read_roles": len({row[1] for row in read_rows if row[1]}),
        "all_grants_basic": all_depths.get("Basic", 0),
        "all_grants_local": all_depths.get("Local", 0),
        "all_grants_deep": all_depths.get("Deep", 0),
        "all_grants_global": all_depths.get("Global", 0),
        "read_grants_basic": read_depths.get("Basic", 0),
        "read_grants_local": read_depths.get("Local", 0),
        "read_grants_deep": read_depths.get("Deep", 0),
        "read_grants_global": read_depths.get("Global", 0),
    }


def _style_sheet(sheet, freeze_panes: str = "A2") -> None:
    sheet.freeze_panes = freeze_panes
    sheet.sheet_view.showGridLines = False
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _style_header(sheet, row: int, columns: int) -> None:
    for column in range(1, columns + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )


def _autosize(sheet, maximum_width: int = 52) -> None:
    for column_index in range(1, sheet.max_column + 1):
        width = max(
            (
                len(str(sheet.cell(row=row, column=column_index).value or ""))
                for row in range(1, sheet.max_row + 1)
            ),
            default=10,
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            maximum_width, max(12, width + 2)
        )


def _build_comparison_sheet(
    workbook: Workbook,
    current: dict[str, Any],
    client: dict[str, Any],
) -> None:
    sheet = workbook.active
    sheet.title = "Comparison"
    sheet.append(["Dataverse Security Scale Comparison"])
    sheet["A1"].font = TITLE_FONT
    sheet.merge_cells("A1:F1")
    sheet.append(
        [
            "Metric",
            "Current Environment",
            "Client Reported",
            "Client Workbook",
            "Scale vs Current",
            "Interpretation / Source",
        ]
    )
    _style_header(sheet, 2, 6)

    metrics = [
        (
            "Active users",
            current["users"]["active"],
            12_500,
            None,
            "Client total is reported in project instructions; workbook has no users.",
        ),
        (
            "Security role instances / reported roles",
            current["security_roles"]["role_instances"],
            450,
            client["exported_roles"],
            "Workbook covers 194 exported definitions; count basis may differ from live BU role instances.",
        ),
        (
            "Unique role names",
            current["security_roles"]["unique_role_names"],
            None,
            client["unique_role_names"],
            "Current value is live; client value is exact names in the workbook.",
        ),
        (
            "Active business units",
            current["business_units"]["active"],
            None,
            None,
            "Client BU hierarchy is not present in the workbook and must be inventoried separately.",
        ),
        (
            "Teams",
            current["teams"]["total"],
            None,
            None,
            "Client team and membership counts are not present in the workbook.",
        ),
        (
            "Field security profiles",
            current["field_security"]["profiles"],
            25,
            None,
            "Client total is reported in project instructions.",
        ),
        (
            "Raw Read privilege grants",
            current["read_privileges"]["raw_role_privilege_entries"],
            None,
            client["read_grants"],
            "Both are role-to-Read-grant rows, but exported role coverage differs.",
        ),
        (
            "Unique Read targets/tables",
            current["read_privileges"]["unique_target_tables"],
            None,
            client["read_unique_targets"],
            "Client workbook contains substantially broader table coverage.",
        ),
        (
            "Deep Read/effective role-table entries",
            current["read_privileges"]["effective_depth_distribution"]["Deep"],
            None,
            client["read_grants_deep"],
            "Current is effective role-table count; workbook is raw role-grant count, so use directionally.",
        ),
        (
            "Maximum generated Deep SQL characters",
            current["deep_predicates"]["maximum_generated_sql_characters"],
            4096,
            None,
            "4,096 is the reported client failure boundary, not a client inventory count.",
        ),
        (
            "Users resolvable to Entra object IDs",
            current["users"]["with_entra_object_id"],
            None,
            None,
            "Only resolvable identities can become OneLake members.",
        ),
    ]

    for row_number, (
        metric,
        current_value,
        reported,
        workbook_value,
        note,
    ) in enumerate(metrics, start=3):
        sheet.append([metric, current_value, reported, workbook_value, None, note])
        sheet.cell(row=row_number, column=2).font = INPUT_FONT
        sheet.cell(row=row_number, column=3).font = INPUT_FONT
        sheet.cell(row=row_number, column=4).font = INPUT_FONT
        sheet.cell(row=row_number, column=5).value = (
            f'=IF(B{row_number}=0,"",IF(ISNUMBER(C{row_number}),'
            f"C{row_number}/B{row_number},IF(ISNUMBER(D{row_number}),"
            f'D{row_number}/B{row_number},"")))'
        )
        sheet.cell(row=row_number, column=5).font = FORMULA_FONT
        sheet.cell(row=row_number, column=5).number_format = "0.0x"

    sheet.append([])
    status_row = sheet.max_row + 1
    sheet.append(
        ["Compile-only smoke-test result", "Configured mapping", "RLS-only diagnostic"]
    )
    _style_header(sheet, status_row, 3)
    configured = current.get("policy_compilation", {})
    rls_only = current.get("policy_compilation_without_cls", {})
    sheet.append(
        [
            "Status",
            f"{configured.get('status', 'not run')}: {configured.get('category', '')}",
            rls_only.get("status", "not run"),
        ]
    )
    sheet.append(
        [
            "Generated policy objects",
            None,
            rls_only.get("generated_policies"),
        ]
    )
    sheet.append(
        [
            "Maximum row rule characters",
            None,
            rls_only.get("maximum_generated_row_rule_characters"),
        ]
    )
    sheet.cell(row=status_row + 1, column=2).fill = WARNING_FILL
    sheet.cell(row=status_row + 1, column=3).fill = GOOD_FILL
    _style_sheet(sheet, "A3")
    sheet["A1"].font = TITLE_FONT
    _style_header(sheet, 2, 6)
    _style_header(sheet, status_row, 3)
    _autosize(sheet)


def _build_flat_sheet(workbook: Workbook, title: str, values: dict[str, Any]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(["Metric", "Value"])
    for metric, value in _flatten(values):
        sheet.append([metric, value])
    _style_sheet(sheet)
    _style_header(sheet, 1, 2)
    _autosize(sheet)


def _build_test_plan_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Smoke Test Plan")
    sheet.append(
        [
            "Phase",
            "Representative fixture",
            "Purpose",
            "Expected evidence",
            "Safety boundary",
        ]
    )
    rows = [
        (
            "1. Unit/in-memory",
            "Synthetic BU graphs at 25, 26, 104, 105, and 160 nodes",
            "Test candidate 1,000-character and observed 4,096-character boundaries.",
            "Exact generated rule lengths and deterministic chunk counts.",
            "No Dataverse or Fabric changes.",
        ),
        (
            "2. Functional Dataverse",
            "4 BUs, 8-12 custom roles, 6-10 test users, owner/access/AAD teams",
            "Cover Basic, Local, Deep, Global, mixed depth, and overlapping memberships.",
            "Policy export matches Dataverse-visible rows for each persona.",
            "Dedicated sandbox; synthetic records only.",
        ),
        (
            "3. CLS composition",
            "3 profiles: uniform, divergent, and no-read grant",
            "Verify CLS grouping and fail-closed behavior.",
            "Configured compile identifies unsupported multi-role RLS+CLS cases.",
            "Do not disable CLS as a production workaround.",
        ),
        (
            "4. Fabric limit probe",
            "One disposable mirrored/Lakehouse table with predicates at boundary lengths",
            "Confirm the actual GA API limit in the target tenant.",
            "PUT acceptance/rejection at N-1, N, and N+1 characters.",
            "Dedicated Fabric item; never the customer item.",
        ),
        (
            "5. Scale simulation",
            "12,500 synthetic principals, 450 role definitions, 25 profile signatures in memory",
            "Measure compilation time, memory, role count, and member/permission chunking.",
            "Capacity report stays below or clearly fails the 1,000-role quota.",
            "Do not create 12,500 real Dataverse users.",
        ),
        (
            "6. Known gaps",
            "RecordFilter mask 16, POA record share, POAA field share",
            "Document expected fail-closed or under-grant behavior.",
            "Operator warnings and explicit unsupported-scenario results.",
            "Never claim exact parity until these are implemented.",
        ),
    ]
    for row in rows:
        sheet.append(row)
    _style_sheet(sheet)
    _style_header(sheet, 1, 5)
    _autosize(sheet)


def _build_do_dont_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Do and Do Not")
    sheet.append(["Disposition", "Action", "Reason"])
    rows = [
        (
            "DO",
            "Use a dedicated Dataverse sandbox and disposable Fabric item.",
            "Prevents access changes in production.",
        ),
        (
            "DO",
            "Import only approved custom role definitions or recreate a representative subset.",
            "Role names/depths are useful; production assignments and data are not required.",
        ),
        (
            "DO",
            "Inventory client BUs, teams, memberships, Entra resolvability, and profile assignments separately.",
            "Those dimensions are absent from the workbook but drive predicates and role counts.",
        ),
        (
            "DO",
            "Run compile-only inventory before every Fabric smoke test.",
            "Catches role capacity and RLS/CLS composition failures without a PUT.",
        ),
        (
            "DO NOT",
            "Clone 12,500 users, production records, or client identities into this tenant.",
            "High operational, privacy, licensing, and access-control risk.",
        ),
        (
            "DO NOT",
            "Run dataverse_test.py against a production Fabric item for diagnostics.",
            "It invokes WeaverAgent.run and can replace data access roles.",
        ),
        (
            "DO NOT",
            "Treat disabling CLS as a fix for the current compile failure.",
            "That can expose secured columns.",
        ),
        (
            "DO NOT",
            "Assume the client workbook is the complete 450-role environment.",
            "It contains 194 exported role definitions and no assignments or BU topology.",
        ),
        (
            "DO NOT",
            "Assume 4,096 is the only current limit without an isolated API probe.",
            "The current REST schema documents no RowConstraint.value maximum.",
        ),
    ]
    for row in rows:
        sheet.append(row)
    _style_sheet(sheet)
    _style_header(sheet, 1, 3)
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=1).fill = (
            GOOD_FILL if sheet.cell(row=row, column=1).value == "DO" else WARNING_FILL
        )
    _autosize(sheet)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--current", required=True)
    parser.add_argument("--client-workbook", required=True)
    parser.add_argument("--output", required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    current_path = Path(args.current)
    client_path = Path(args.client_workbook)
    output_path = Path(args.output)

    current = json.loads(current_path.read_text(encoding="utf-8"))
    client = _client_workbook_summary(client_path)

    workbook = Workbook()
    _build_comparison_sheet(workbook, current, client)
    _build_flat_sheet(workbook, "Current Inventory", current)
    _build_flat_sheet(workbook, "Client Workbook", client)
    _build_test_plan_sheet(workbook)
    _build_do_dont_sheet(workbook)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.font == Font():
                    cell.font = BODY_FONT

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


if __name__ == "__main__":
    main()
