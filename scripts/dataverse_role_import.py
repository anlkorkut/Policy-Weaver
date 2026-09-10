"""Plan, apply, verify, or roll back Dataverse security-role imports.

The workbook is matched to the target environment by exact privilege name.
Dry-run planning is the default. Existing roles are never changed, and every
created role is recorded in a rollback manifest.
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from policyweaver.core.auth import ServicePrincipal
from policyweaver.core.conf import Configuration
from policyweaver.plugins.dataverse.model import DataverseSourceMap


DEPTH_RANK = {"Basic": 0, "Local": 1, "Deep": 2, "Global": 3}
DEPTH_BY_MASK = {1: "Basic", 2: "Local", 4: "Deep", 8: "Global"}
CAPABILITY_BY_DEPTH = {
    "Basic": "canbebasic",
    "Local": "canbelocal",
    "Deep": "canbedeep",
    "Global": "canbeglobal",
}
ROLE_ID_PATTERN = re.compile(r"roles\(([0-9a-fA-F-]{36})\)")


class DataverseRoleImporter:
    def __init__(
        self,
        config: DataverseSourceMap,
        organization_id: str,
        environment_id: str,
    ) -> None:
        if not config.dataverse or not config.dataverse.environment_url:
            raise ValueError("dataverse.environment_url configuration is required")
        if not config.service_principal:
            raise ValueError("service_principal configuration is required")

        self.config = config
        self.expected_organization_id = organization_id.lower()
        self.expected_environment_id = environment_id.lower()
        self.base_url = config.dataverse.environment_url.rstrip("/")
        self.api_url = f"{self.base_url}/api/data/v9.2"
        self.scope = f"{self.base_url}/.default"
        self.logger = logging.getLogger("DATAVERSE_ROLE_IMPORT")

        retry = Retry(
            total=5,
            connect=5,
            read=5,
            backoff_factor=1.0,
            status_forcelist=[429, 500, 502, 503, 504],
            allowed_methods=frozenset(["GET"]),
        )
        self.session = requests.Session()
        self.session.mount("https://", HTTPAdapter(max_retries=retry))
        self.timeout = (10, 120)

    @property
    def headers(self) -> dict[str, str]:
        token = ServicePrincipal.Credential.get_token(self.scope).token
        return {
            "Authorization": f"Bearer {token}",
            "Accept": "application/json",
            "Content-Type": "application/json; charset=utf-8",
            "OData-MaxVersion": "4.0",
            "OData-Version": "4.0",
            "If-None-Match": "null",
            "Prefer": "odata.maxpagesize=5000",
        }

    def get_json(self, url: str) -> dict[str, Any]:
        response = self.session.get(url, headers=self.headers, timeout=self.timeout)
        response.raise_for_status()
        return response.json()

    def get_paged(self, url: str) -> list[dict[str, Any]]:
        values: list[dict[str, Any]] = []
        while url:
            data = self.get_json(url)
            values.extend(data.get("value", []))
            url = data.get("@odata.nextLink")
        return values

    def validate_target(self) -> dict[str, str]:
        who_am_i = self.get_json(f"{self.api_url}/WhoAmI")
        actual_organization_id = str(who_am_i.get("OrganizationId", "")).lower()
        if actual_organization_id != self.expected_organization_id:
            raise ValueError(
                "Organization ID mismatch: authenticated Dataverse organization "
                "does not match --organization-id"
            )

        organization = self.get_json(
            f"{self.api_url}/RetrieveCurrentOrganization(AccessType=@p1)"
            "?@p1=Microsoft.Dynamics.CRM.EndpointAccessType'Default'"
        ).get("Detail", {})
        actual_environment_id = str(organization.get("EnvironmentId", "")).lower()
        normalized_actual_environment_id = actual_environment_id.removeprefix(
            "default-"
        )
        if self.expected_environment_id not in {
            actual_environment_id,
            normalized_actual_environment_id,
        }:
            raise ValueError(
                "Environment ID mismatch: authenticated Dataverse environment "
                "does not match --environment-id"
            )

        return {
            "organization_id": actual_organization_id,
            "environment_id": actual_environment_id,
            "organization_version": str(organization.get("OrganizationVersion", "")),
        }

    def root_business_unit(self) -> dict[str, Any]:
        business_units = self.get_paged(
            f"{self.api_url}/businessunits"
            "?$select=businessunitid,name,_parentbusinessunitid_value,isdisabled"
            "&$filter=isdisabled eq false"
        )
        roots = [
            business_unit
            for business_unit in business_units
            if not business_unit.get("_parentbusinessunitid_value")
        ]
        if len(roots) != 1:
            raise ValueError(
                f"Expected one active root business unit, found {len(roots)}"
            )
        return roots[0]

    def existing_roles(self) -> list[dict[str, Any]]:
        return self.get_paged(
            f"{self.api_url}/roles"
            "?$select=roleid,name,_businessunitid_value,"
            "_parentroleid_value,_parentrootroleid_value,componentstate"
            "&$filter=componentstate eq 0"
        )

    def privileges(self) -> list[dict[str, Any]]:
        return self.get_paged(
            f"{self.api_url}/privileges"
            "?$select=privilegeid,name,canbebasic,canbelocal,canbedeep,canbeglobal"
        )

    def role_privilege_names(self, role_id: str) -> set[str]:
        values = self.get_paged(
            f"{self.api_url}/roles({role_id})/roleprivileges_association?$select=name"
        )
        return {str(value.get("name")) for value in values if value.get("name")}

    def role_privilege_depths(self, role_id: str) -> dict[str, str]:
        values = self.get_paged(
            f"{self.api_url}/roleprivilegescollection"
            "?$select=privilegeid,privilegedepthmask"
            f"&$filter=componentstate eq 0 and roleid eq {role_id}"
        )
        return {
            str(value["privilegeid"]).lower(): DEPTH_BY_MASK.get(
                value.get("privilegedepthmask"), "Unknown"
            )
            for value in values
            if value.get("privilegeid")
        }

    def all_role_privilege_depths(self) -> dict[str, dict[str, str]]:
        values = self.get_paged(
            f"{self.api_url}/roleprivilegescollection"
            "?$select=roleid,privilegeid,privilegedepthmask"
            "&$filter=componentstate eq 0"
        )
        result: dict[str, dict[str, str]] = defaultdict(dict)
        for value in values:
            role_id = str(value.get("roleid", "")).lower()
            privilege_id = str(value.get("privilegeid", "")).lower()
            if not role_id or not privilege_id:
                continue
            result[role_id][privilege_id] = DEPTH_BY_MASK.get(
                value.get("privilegedepthmask"), "Unknown"
            )
        return dict(result)

    def create_role(self, name: str, business_unit_id: str) -> str:
        response = self.session.post(
            f"{self.api_url}/roles",
            headers=self.headers,
            json={
                "businessunitid@odata.bind": f"businessunits({business_unit_id})",
                "name": name,
            },
            timeout=self.timeout,
        )
        response.raise_for_status()
        entity_id = response.headers.get("OData-EntityId", "")
        match = ROLE_ID_PATTERN.search(entity_id)
        if not match:
            raise RuntimeError("Dataverse role creation did not return OData-EntityId")
        return match.group(1).lower()

    def add_privileges(self, role_id: str, privileges: list[dict[str, Any]]) -> None:
        response = self.session.post(
            f"{self.api_url}/roles({role_id})/Microsoft.Dynamics.CRM.AddPrivilegesRole",
            headers=self.headers,
            json={"Privileges": privileges},
            timeout=self.timeout,
        )
        response.raise_for_status()

    def replace_privileges(
        self, role_id: str, privileges: list[dict[str, Any]]
    ) -> None:
        response = self.session.post(
            f"{self.api_url}/roles({role_id})/"
            "Microsoft.Dynamics.CRM.ReplacePrivilegesRole",
            headers=self.headers,
            json={"Privileges": privileges},
            timeout=self.timeout,
        )
        response.raise_for_status()

    def delete_role(self, role_id: str) -> None:
        response = self.session.delete(
            f"{self.api_url}/roles({role_id})",
            headers=self.headers,
            timeout=self.timeout,
        )
        if response.status_code == 404:
            return
        response.raise_for_status()


def read_workbook_roles(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    roles_sheet = workbook["Roles"]
    privileges_sheet = workbook["Role_Privileges"]

    role_name_by_index: dict[Any, str] = {}
    source_role_id_by_index: dict[Any, str] = {}
    for row in roles_sheet.iter_rows(min_row=2, values_only=True):
        role_index = row[0]
        role_name = str(row[2] or row[1] or "").strip()
        if role_index is None or not role_name:
            continue
        role_name_by_index[role_index] = role_name
        source_role_id_by_index[role_index] = str(row[3] or "")

    privileges_by_role: dict[Any, dict[str, str]] = defaultdict(dict)
    for row in privileges_sheet.iter_rows(min_row=2, values_only=True):
        role_index = row[0]
        privilege_name = str(row[4] or "").strip()
        depth = str(row[7] or "").strip().title()
        if role_index not in role_name_by_index or not privilege_name:
            continue
        if depth not in DEPTH_RANK:
            continue
        current_depth = privileges_by_role[role_index].get(privilege_name)
        if current_depth is None or DEPTH_RANK[depth] >= DEPTH_RANK[current_depth]:
            privileges_by_role[role_index][privilege_name] = depth

    logical_roles: dict[str, dict[str, Any]] = {}
    for role_index, role_name in sorted(
        role_name_by_index.items(), key=lambda item: str(item[1]).casefold()
    ):
        normalized_name = role_name.casefold()
        logical_role = logical_roles.setdefault(
            normalized_name,
            {
                "source_role_index": role_index,
                "source_role_indices": [],
                "source_role_id": source_role_id_by_index.get(role_index, ""),
                "source_role_ids": [],
                "name": role_name,
                "privileges": {},
            },
        )
        logical_role["source_role_indices"].append(role_index)
        source_role_id = source_role_id_by_index.get(role_index, "")
        if source_role_id:
            logical_role["source_role_ids"].append(source_role_id)

        for privilege_name, depth in privileges_by_role.get(role_index, {}).items():
            current_depth = logical_role["privileges"].get(privilege_name)
            if current_depth is None or DEPTH_RANK[depth] >= DEPTH_RANK[current_depth]:
                logical_role["privileges"][privilege_name] = depth

    return list(logical_roles.values())


def build_plan(
    importer: DataverseRoleImporter,
    source_roles: list[dict[str, Any]],
) -> dict[str, Any]:
    target = importer.validate_target()
    root_business_unit = importer.root_business_unit()
    existing_roles = importer.existing_roles()
    target_privileges = importer.privileges()

    existing_names = {
        str(role.get("name", "")).casefold(): role
        for role in existing_roles
        if role.get("name")
    }
    privilege_by_name = {
        str(privilege.get("name")): privilege
        for privilege in target_privileges
        if privilege.get("name")
    }

    role_plans: list[dict[str, Any]] = []
    missing_privilege_names: Counter[str] = Counter()
    unsupported_depths: Counter[str] = Counter()
    for source_role in source_roles:
        collision = existing_names.get(source_role["name"].casefold())
        matched: list[dict[str, Any]] = []
        missing: list[str] = []
        unsupported: list[dict[str, str]] = []

        for privilege_name, depth in source_role["privileges"].items():
            target_privilege = privilege_by_name.get(privilege_name)
            if not target_privilege:
                missing.append(privilege_name)
                missing_privilege_names[privilege_name] += 1
                continue
            capability = CAPABILITY_BY_DEPTH[depth]
            if not bool(target_privilege.get(capability)):
                unsupported.append({"name": privilege_name, "depth": depth})
                unsupported_depths[depth] += 1
                continue
            matched.append(
                {
                    "name": privilege_name,
                    "depth": depth,
                    "privilege_id": target_privilege["privilegeid"],
                }
            )

        role_plans.append(
            {
                "source_role_index": source_role["source_role_index"],
                "source_role_id": source_role["source_role_id"],
                "name": source_role["name"],
                "status": "collision" if collision else "ready",
                "existing_role_id": collision.get("roleid") if collision else None,
                "source_privilege_count": len(source_role["privileges"]),
                "matched_privileges": matched,
                "missing_privileges": sorted(missing),
                "unsupported_depths": sorted(
                    unsupported, key=lambda value: (value["name"], value["depth"])
                ),
            }
        )

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "target": target,
        "root_business_unit": {
            "business_unit_id": root_business_unit["businessunitid"],
            "name": root_business_unit.get("name"),
        },
        "summary": {
            "source_roles": len(source_roles),
            "existing_target_roles": len(existing_roles),
            "role_name_collisions": sum(
                role["status"] == "collision" for role in role_plans
            ),
            "roles_ready_to_create": sum(
                role["status"] == "ready" for role in role_plans
            ),
            "source_privilege_assignments": sum(
                role["source_privilege_count"] for role in role_plans
            ),
            "matched_privilege_assignments": sum(
                len(role["matched_privileges"]) for role in role_plans
            ),
            "missing_privilege_assignments": sum(
                len(role["missing_privileges"]) for role in role_plans
            ),
            "unsupported_depth_assignments": sum(
                len(role["unsupported_depths"]) for role in role_plans
            ),
            "unique_missing_privileges": len(missing_privilege_names),
            "unsupported_depths": dict(sorted(unsupported_depths.items())),
        },
        "roles": role_plans,
    }


def write_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


def configure(config_path: Path) -> DataverseSourceMap:
    config = DataverseSourceMap.from_yaml(str(config_path))
    Configuration.configure_environment(config)
    if not config.service_principal:
        raise ValueError("service_principal configuration is required")
    ServicePrincipal.initialize(
        tenant_id=config.service_principal.tenant_id,
        client_id=config.service_principal.client_id,
        client_secret=config.service_principal.client_secret,
    )
    if config.dataverse and config.dataverse.environment_url:
        os.environ["DATAVERSE_ENVIRONMENT_URL"] = config.dataverse.environment_url
    return config


def apply_plan(
    importer: DataverseRoleImporter,
    plan: dict[str, Any],
    manifest_path: Path,
    chunk_size: int,
) -> dict[str, Any]:
    business_unit_id = plan["root_business_unit"]["business_unit_id"]
    manifest: dict[str, Any] = {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "target": plan["target"],
        "root_business_unit": plan["root_business_unit"],
        "created_roles": [],
        "failed_roles": [],
    }
    write_json(manifest_path, manifest)

    for role_plan in plan["roles"]:
        if role_plan["status"] != "ready":
            continue

        role_id: str | None = None
        try:
            role_id = importer.create_role(role_plan["name"], business_unit_id)
            manifest_entry = {
                "role_id": role_id,
                "name": role_plan["name"],
                "status": "created",
                "applied_privileges": 0,
            }
            manifest["created_roles"].append(manifest_entry)
            write_json(manifest_path, manifest)

            default_privileges = importer.role_privilege_names(role_id)
            additions = [
                privilege
                for privilege in role_plan["matched_privileges"]
                if privilege["name"] not in default_privileges
            ]
            for start in range(0, len(additions), chunk_size):
                chunk = additions[start : start + chunk_size]
                payload = [
                    {
                        "Depth": privilege["depth"],
                        "PrivilegeId": privilege["privilege_id"],
                        "BusinessUnitId": business_unit_id,
                        "PrivilegeName": privilege["name"],
                    }
                    for privilege in chunk
                ]
                if payload:
                    importer.add_privileges(role_id, payload)
                    manifest_entry["applied_privileges"] += len(payload)
                    write_json(manifest_path, manifest)

            actual_privileges = importer.role_privilege_names(role_id)
            requested_names = {
                privilege["name"] for privilege in role_plan["matched_privileges"]
            }
            manifest_entry["verified_requested_privileges"] = len(
                requested_names.intersection(actual_privileges)
            )
            manifest_entry["status"] = "verified"
            write_json(manifest_path, manifest)
        except Exception as error:
            failure = {
                "name": role_plan["name"],
                "exception_type": type(error).__name__,
            }
            if isinstance(error, requests.HTTPError) and error.response is not None:
                failure["http_status"] = error.response.status_code
                try:
                    dataverse_error = error.response.json().get("error", {})
                    failure["dataverse_error_code"] = dataverse_error.get("code")
                    failure["dataverse_error_message"] = dataverse_error.get("message")
                except requests.JSONDecodeError:
                    failure["dataverse_error_code"] = None
            if role_id:
                try:
                    importer.delete_role(role_id)
                    failure["partial_role_deleted"] = True
                    manifest["created_roles"] = [
                        role
                        for role in manifest["created_roles"]
                        if role["role_id"] != role_id
                    ]
                except Exception:
                    failure["partial_role_deleted"] = False
            manifest["failed_roles"].append(failure)
            write_json(manifest_path, manifest)

    manifest["completed_at"] = datetime.now(timezone.utc).isoformat()
    manifest["summary"] = {
        "created_and_verified": sum(
            role["status"] == "verified" for role in manifest["created_roles"]
        ),
        "failed": len(manifest["failed_roles"]),
        "applied_privileges": sum(
            role["applied_privileges"] for role in manifest["created_roles"]
        ),
    }
    write_json(manifest_path, manifest)
    return manifest


def rollback(
    importer: DataverseRoleImporter,
    manifest: dict[str, Any],
    output_path: Path,
) -> dict[str, Any]:
    importer.validate_target()
    results = []
    for role in reversed(manifest.get("created_roles", [])):
        try:
            importer.delete_role(role["role_id"])
            results.append({"role_id": role["role_id"], "status": "deleted"})
        except Exception as error:
            results.append(
                {
                    "role_id": role["role_id"],
                    "status": "failed",
                    "exception_type": type(error).__name__,
                }
            )
    report = {
        "rolled_back_at": datetime.now(timezone.utc).isoformat(),
        "target": manifest.get("target"),
        "results": results,
    }
    write_json(output_path, report)
    return report


def verify_import(
    importer: DataverseRoleImporter,
    plan: dict[str, Any],
    manifest: dict[str, Any],
    output_path: Path,
    repair_privileges: bool = False,
) -> dict[str, Any]:
    importer.validate_target()
    plan_by_name = {role["name"]: role for role in plan.get("roles", [])}
    results = []
    for created_role in manifest.get("created_roles", []):
        role_plan = plan_by_name.get(created_role["name"])
        if not role_plan:
            results.append(
                {
                    "role_id": created_role["role_id"],
                    "name": created_role["name"],
                    "status": "missing_from_plan",
                }
            )
            continue

        expected = {
            privilege["privilege_id"].lower(): privilege["depth"]
            for privilege in role_plan["matched_privileges"]
        }
        try:
            actual_depths = importer.role_privilege_depths(created_role["role_id"])

            def mismatch_counts() -> tuple[int, int, int]:
                missing = sum(
                    privilege_id not in actual_depths for privilege_id in expected
                )
                wrong_depth = sum(
                    privilege_id in actual_depths
                    and actual_depths[privilege_id] != expected_depth
                    for privilege_id, expected_depth in expected.items()
                )
                unexpected = len(set(actual_depths).difference(expected))
                return missing, wrong_depth, unexpected

            missing_count, depth_mismatch_count, unexpected_count = mismatch_counts()
            repaired = False
            if repair_privileges and any(
                (missing_count, depth_mismatch_count, unexpected_count)
            ):
                business_unit_id = plan["root_business_unit"]["business_unit_id"]
                payload = [
                    {
                        "Depth": privilege["depth"],
                        "PrivilegeId": privilege["privilege_id"],
                        "BusinessUnitId": business_unit_id,
                        "PrivilegeName": privilege["name"],
                    }
                    for privilege in role_plan["matched_privileges"]
                ]
                importer.replace_privileges(created_role["role_id"], payload)
                repaired = True
                actual_depths = importer.role_privilege_depths(created_role["role_id"])
                missing_count, depth_mismatch_count, unexpected_count = (
                    mismatch_counts()
                )
        except Exception as error:
            result = {
                "role_id": created_role["role_id"],
                "name": created_role["name"],
                "status": "verification_failed",
                "exception_type": type(error).__name__,
            }
            if isinstance(error, requests.HTTPError) and error.response is not None:
                result["http_status"] = error.response.status_code
                try:
                    dataverse_error = error.response.json().get("error", {})
                    result["dataverse_error_code"] = dataverse_error.get("code")
                    result["dataverse_error_message"] = dataverse_error.get("message")
                except requests.JSONDecodeError:
                    result["dataverse_error_code"] = None
            results.append(result)
            continue

        results.append(
            {
                "role_id": created_role["role_id"],
                "name": created_role["name"],
                "status": (
                    "verified"
                    if missing_count == 0
                    and depth_mismatch_count == 0
                    and unexpected_count == 0
                    else "mismatch"
                ),
                "expected_privileges": len(expected),
                "missing_privileges": missing_count,
                "depth_mismatches": depth_mismatch_count,
                "unexpected_privileges": unexpected_count,
                "repaired": repaired,
            }
        )

    report = {
        "verified_at": datetime.now(timezone.utc).isoformat(),
        "target": manifest.get("target"),
        "summary": {
            "roles_checked": len(results),
            "verified": sum(result["status"] == "verified" for result in results),
            "mismatched": sum(result["status"] == "mismatch" for result in results),
            "missing_from_plan": sum(
                result["status"] == "missing_from_plan" for result in results
            ),
            "verification_failed": sum(
                result["status"] == "verification_failed" for result in results
            ),
            "missing_privileges": sum(
                result.get("missing_privileges", 0) for result in results
            ),
            "depth_mismatches": sum(
                result.get("depth_mismatches", 0) for result in results
            ),
            "unexpected_privileges": sum(
                result.get("unexpected_privileges", 0) for result in results
            ),
            "repair_requested": repair_privileges,
            "roles_repaired": sum(result.get("repaired", False) for result in results),
        },
        "results": results,
    }
    write_json(output_path, report)
    return report


def verify_all_instances(
    importer: DataverseRoleImporter,
    plan: dict[str, Any],
    manifest: dict[str, Any],
    output_path: Path,
) -> dict[str, Any]:
    importer.validate_target()
    imported_names = {
        role["name"].casefold(): role["name"]
        for role in manifest.get("created_roles", [])
    }
    plan_by_name = {role["name"].casefold(): role for role in plan.get("roles", [])}
    instances_by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for role in importer.existing_roles():
        name = str(role.get("name", ""))
        if name.casefold() in imported_names:
            instances_by_name[name.casefold()].append(role)
    depths_by_role_id = importer.all_role_privilege_depths()

    results = []
    for normalized_name, role_name in sorted(imported_names.items()):
        role_plan = plan_by_name.get(normalized_name)
        if not role_plan:
            results.append({"name": role_name, "status": "missing_from_plan"})
            continue
        expected = {
            privilege["privilege_id"].lower(): privilege["depth"]
            for privilege in role_plan["matched_privileges"]
        }
        instances = []
        for role in instances_by_name.get(normalized_name, []):
            role_id = str(role.get("roleid", "")).lower()
            parent_root_role_id = str(
                role.get("_parentrootroleid_value", "") or ""
            ).lower()
            privilege_source_role_id = (
                role_id if role_id in depths_by_role_id else parent_root_role_id
            )
            actual = depths_by_role_id.get(privilege_source_role_id, {})
            missing = sum(privilege_id not in actual for privilege_id in expected)
            wrong_depth = sum(
                privilege_id in actual and actual[privilege_id] != expected_depth
                for privilege_id, expected_depth in expected.items()
            )
            unexpected = len(set(actual).difference(expected))
            instances.append(
                {
                    "role_id": role_id,
                    "business_unit_id": role.get("_businessunitid_value"),
                    "privilege_source_role_id": privilege_source_role_id,
                    "inherits_from_root": privilege_source_role_id != role_id,
                    "status": (
                        "verified"
                        if missing == 0 and wrong_depth == 0 and unexpected == 0
                        else "mismatch"
                    ),
                    "missing_privileges": missing,
                    "depth_mismatches": wrong_depth,
                    "unexpected_privileges": unexpected,
                }
            )
        results.append(
            {
                "name": role_name,
                "status": (
                    "verified"
                    if instances
                    and all(instance["status"] == "verified" for instance in instances)
                    else "mismatch"
                ),
                "instances": instances,
            }
        )

    instance_results = [
        instance for result in results for instance in result.get("instances", [])
    ]
    report = {
        "verified_at": datetime.now(timezone.utc).isoformat(),
        "target": manifest.get("target"),
        "summary": {
            "logical_roles_checked": len(results),
            "logical_roles_verified": sum(
                result["status"] == "verified" for result in results
            ),
            "logical_roles_mismatched": sum(
                result["status"] == "mismatch" for result in results
            ),
            "instances_checked": len(instance_results),
            "instances_verified": sum(
                instance["status"] == "verified" for instance in instance_results
            ),
            "instances_mismatched": sum(
                instance["status"] == "mismatch" for instance in instance_results
            ),
            "missing_privileges": sum(
                instance["missing_privileges"] for instance in instance_results
            ),
            "depth_mismatches": sum(
                instance["depth_mismatches"] for instance in instance_results
            ),
            "unexpected_privileges": sum(
                instance["unexpected_privileges"] for instance in instance_results
            ),
        },
        "results": results,
    }
    write_json(output_path, report)
    return report


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default="configdataverse.yaml")
    parser.add_argument(
        "--workbook", default="Dataverse_Security_Roles_Consolidated.xlsx"
    )
    parser.add_argument("--organization-id", required=True)
    parser.add_argument("--environment-id", required=True)
    parser.add_argument(
        "--role-name",
        action="append",
        help="Import only an exact role name. Repeat to select multiple roles.",
    )
    parser.add_argument(
        "--max-roles",
        type=int,
        help="Limit the selected source roles after sorting. Intended for canaries.",
    )
    parser.add_argument(
        "--plan-output", default="reports/dataverse-role-import-plan.json"
    )
    parser.add_argument("--apply", action="store_true")
    parser.add_argument(
        "--allow-partial",
        action="store_true",
        help="Allow --apply when client privileges are unavailable in the target.",
    )
    parser.add_argument(
        "--manifest", default="reports/dataverse-role-import-manifest.json"
    )
    parser.add_argument("--rollback-manifest")
    parser.add_argument(
        "--rollback-output", default="reports/dataverse-role-import-rollback.json"
    )
    parser.add_argument("--verify-manifest")
    parser.add_argument(
        "--verify-plan", default="reports/dataverse-role-import-plan-applied.json"
    )
    parser.add_argument(
        "--verify-output", default="reports/dataverse-role-import-verification.json"
    )
    parser.add_argument("--repair-privileges", action="store_true")
    parser.add_argument("--verify-all-instances", action="store_true")
    parser.add_argument(
        "--verify-instances-output",
        default="reports/dataverse-role-import-instance-verification.json",
    )
    parser.add_argument("--chunk-size", type=int, default=100)
    parser.add_argument(
        "--log-level",
        default="INFO",
        choices=["DEBUG", "INFO", "WARNING", "ERROR"],
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    logging.basicConfig(level=getattr(logging, args.log_level))
    config = configure(Path(args.config))
    importer = DataverseRoleImporter(
        config=config,
        organization_id=args.organization_id,
        environment_id=args.environment_id,
    )

    if args.rollback_manifest:
        manifest = json.loads(Path(args.rollback_manifest).read_text(encoding="utf-8"))
        report = rollback(importer, manifest, Path(args.rollback_output))
        print(
            json.dumps(
                {"rollback": Counter(r["status"] for r in report["results"])}, indent=2
            )
        )
        return

    if args.verify_manifest:
        manifest = json.loads(Path(args.verify_manifest).read_text(encoding="utf-8"))
        plan = json.loads(Path(args.verify_plan).read_text(encoding="utf-8"))
        if args.role_name:
            selected_names = {name.casefold() for name in args.role_name}
            manifest = {
                **manifest,
                "created_roles": [
                    role
                    for role in manifest.get("created_roles", [])
                    if role["name"].casefold() in selected_names
                ],
            }
            plan = {
                **plan,
                "roles": [
                    role
                    for role in plan.get("roles", [])
                    if role["name"].casefold() in selected_names
                ],
            }
        if args.verify_all_instances:
            report = verify_all_instances(
                importer, plan, manifest, Path(args.verify_instances_output)
            )
        else:
            report = verify_import(
                importer,
                plan,
                manifest,
                Path(args.verify_output),
                repair_privileges=args.repair_privileges,
            )
        print(json.dumps(report["summary"], indent=2, sort_keys=True))
        return

    source_roles = read_workbook_roles(Path(args.workbook))
    if args.role_name:
        requested_names = {name.casefold(): name for name in args.role_name}
        source_roles = [
            role for role in source_roles if role["name"].casefold() in requested_names
        ]
        found_names = {role["name"].casefold() for role in source_roles}
        missing_names = sorted(
            original_name
            for normalized_name, original_name in requested_names.items()
            if normalized_name not in found_names
        )
        if missing_names:
            raise ValueError(
                "Requested role names not found in workbook: "
                + ", ".join(missing_names)
            )
    if args.max_roles is not None:
        if args.max_roles < 1:
            raise ValueError("--max-roles must be a positive integer")
        source_roles = source_roles[: args.max_roles]
    plan = build_plan(importer, source_roles)
    write_json(Path(args.plan_output), plan)
    print(json.dumps(plan["summary"], indent=2, sort_keys=True))

    if not args.apply:
        return
    if plan["summary"]["missing_privilege_assignments"] > 0 and not args.allow_partial:
        raise ValueError(
            "Import is partial because target privileges are missing. "
            "Review the plan and pass --allow-partial to proceed."
        )

    manifest = apply_plan(
        importer=importer,
        plan=plan,
        manifest_path=Path(args.manifest),
        chunk_size=args.chunk_size,
    )
    print(json.dumps(manifest["summary"], indent=2, sort_keys=True))


if __name__ == "__main__":
    try:
        main()
    except requests.HTTPError as error:
        response = error.response
        status_code = response.status_code if response is not None else "unknown"
        print(f"Dataverse HTTP error: {status_code}", file=sys.stderr)
        raise
